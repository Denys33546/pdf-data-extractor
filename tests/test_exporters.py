import csv
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from pdf_extractor.exporters import (
    INVOICE_COLUMNS,
    ITEM_COLUMNS,
    UNKNOWN_VENDOR,
    invoice_rows,
    item_rows,
    summary_metrics,
    vendor_totals,
    write_csv,
    write_excel,
)
from pdf_extractor.extractor import ExtractedInvoice, ExtractedLine, extract_folder
from pdf_extractor.sample_data import generate_samples


@pytest.fixture(scope="module")
def invoices(tmp_path_factory):
    folder = tmp_path_factory.mktemp("invoices")
    generate_samples(folder, count=5)
    return extract_folder(folder)


def make_invoice(**overrides):
    values = {
        "source": "a.pdf",
        "number": "INV-1",
        "issue_date": date(2026, 1, 2),
        "vendor": "Acme",
        "customer": "Buyer",
        "items": (ExtractedLine("Widget", 2, Decimal("10.00"), Decimal("20.00")),),
        "subtotal": Decimal("20.00"),
        "tax": Decimal("2.60"),
        "total": Decimal("22.60"),
        "warnings": (),
    }
    values.update(overrides)
    return ExtractedInvoice(**values)


# ---------- plain data ----------


def test_invoice_rows_one_per_invoice(invoices):
    rows = invoice_rows(invoices)
    assert len(rows) == 5
    assert rows[0][1] == invoices[0].number
    assert rows[0][7] == invoices[0].total
    assert {row[8] for row in rows} == {"OK"}


def test_item_rows_flatten_all_items(invoices):
    rows = item_rows(invoices)
    assert len(rows) == sum(len(inv.items) for inv in invoices)
    assert rows[0][4] == invoices[0].items[0].description
    assert all(len(row) == len(ITEM_COLUMNS) for row in rows)


def test_status_and_warnings_columns():
    bad = make_invoice(warnings=("Total not found", "Date not found"))
    row = invoice_rows([bad])[0]
    assert row[8] == "CHECK"
    assert row[9] == "Total not found; Date not found"


def test_vendor_totals_sorted_by_amount():
    invoices = [
        make_invoice(vendor="Small", total=Decimal("10.00")),
        make_invoice(vendor="Big", total=Decimal("500.00")),
        make_invoice(vendor="Big", total=Decimal("100.00")),
        make_invoice(vendor=None, total=None),
    ]
    assert vendor_totals(invoices) == [
        ("Big", 2, Decimal("600.00")),
        ("Small", 1, Decimal("10.00")),
        (UNKNOWN_VENDOR, 1, Decimal("0.00")),
    ]


def test_summary_metrics(invoices):
    metrics = dict(summary_metrics([*invoices, make_invoice(warnings=("x",), total=None)]))
    assert metrics["Invoices processed"] == 6
    assert metrics["All checks passed"] == 5
    assert metrics["Need review"] == 1
    assert metrics["Total of all invoices"] == sum((i.total for i in invoices), Decimal("0.00"))


# ---------- Excel ----------


def test_excel_has_expected_sheets_and_data(invoices, tmp_path):
    path = write_excel(invoices, tmp_path / "report.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Summary", "Invoices", "Items"]

    sheet = workbook["Invoices"]
    assert [cell.value for cell in sheet[1]] == INVOICE_COLUMNS
    assert sheet.max_row == 6  # header + 5 invoices
    assert sheet["B2"].value == invoices[0].number
    assert sheet["C2"].value.date() == invoices[0].issue_date
    assert float(sheet["H2"].value) == float(invoices[0].total)
    assert sheet["H2"].number_format == "#,##0.00"
    assert sheet["C2"].number_format == "yyyy-mm-dd"
    assert sheet.freeze_panes == "A2"

    items = workbook["Items"]
    assert items.max_row == 1 + sum(len(inv.items) for inv in invoices)


def test_excel_summary_sheet(invoices, tmp_path):
    workbook = load_workbook(write_excel(invoices, tmp_path / "report.xlsx"))
    summary = workbook["Summary"]
    assert summary["A1"].value == "Invoices processed"
    assert summary["B1"].value == 5
    labels = [row[0].value for row in summary.iter_rows(min_row=1, max_row=summary.max_row)]
    assert "Vendor" in labels


def test_excel_marks_rows_that_need_review(tmp_path):
    bad = make_invoice(warnings=("Total not found",))
    good = make_invoice(source="b.pdf")
    workbook = load_workbook(write_excel([good, bad], tmp_path / "report.xlsx"))
    sheet = workbook["Invoices"]
    assert sheet["I2"].value == "OK"
    assert sheet["I3"].value == "CHECK"
    assert sheet["A3"].fill.fgColor.rgb.endswith("FDE2E2")
    assert not sheet["A2"].fill.fgColor.rgb.endswith("FDE2E2")


def test_excel_keeps_formula_like_text_as_text(tmp_path):
    invoice = make_invoice(
        vendor='=HYPERLINK("http://evil")',
        items=(ExtractedLine("=1+1", 1, Decimal("1.00"), Decimal("1.00")),),
    )
    workbook = load_workbook(write_excel([invoice], tmp_path / "report.xlsx"))
    vendor_cell = workbook["Invoices"]["D2"]
    item_cell = workbook["Items"]["E2"]
    assert vendor_cell.data_type == "s" and vendor_cell.value.startswith("=")
    assert item_cell.data_type == "s" and item_cell.value == "=1+1"


def test_excel_with_no_invoices_still_has_headers(tmp_path):
    workbook = load_workbook(write_excel([], tmp_path / "report.xlsx"))
    assert [c.value for c in workbook["Invoices"][1]] == INVOICE_COLUMNS
    assert workbook["Summary"]["B1"].value == 0


def test_excel_creates_missing_folders(invoices, tmp_path):
    path = write_excel(invoices, tmp_path / "deep" / "er" / "report.xlsx")
    assert path.exists()


# ---------- CSV ----------


def read_csv(path):
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.reader(handle))


def test_csv_writes_invoices_and_items_files(invoices, tmp_path):
    paths = write_csv(invoices, tmp_path / "out" / "report.csv")
    assert [p.name for p in paths] == ["report_invoices.csv", "report_items.csv"]

    invoice_table = read_csv(paths[0])
    assert invoice_table[0] == INVOICE_COLUMNS
    assert len(invoice_table) == 6
    assert invoice_table[1][1] == invoices[0].number
    assert invoice_table[1][2] == invoices[0].issue_date.isoformat()
    assert invoice_table[1][7] == f"{invoices[0].total:.2f}"

    item_table = read_csv(paths[1])
    assert item_table[0] == ITEM_COLUMNS
    assert len(item_table) == 1 + sum(len(inv.items) for inv in invoices)


def test_csv_empty_values_are_blank(tmp_path):
    invoice = make_invoice(number=None, issue_date=None, total=None, warnings=("x",))
    paths = write_csv([invoice], tmp_path / "report.csv")
    row = read_csv(paths[0])[1]
    assert row[1] == "" and row[2] == "" and row[7] == ""


def test_csv_neutralises_formulas(tmp_path):
    invoice = make_invoice(
        items=(ExtractedLine("=1+1", 1, Decimal("1.00"), Decimal("1.00")),), vendor="@evil"
    )
    paths = write_csv([invoice], tmp_path / "report.csv")
    assert read_csv(paths[1])[1][4] == "'=1+1"
    assert read_csv(paths[0])[1][3] == "'@evil"


def test_csv_starts_with_utf8_bom(invoices, tmp_path):
    paths = write_csv(invoices, tmp_path / "report.csv")
    assert paths[0].read_bytes().startswith(b"\xef\xbb\xbf")


def test_excel_tables_print_landscape_one_page_wide(invoices, tmp_path):
    workbook = load_workbook(write_excel(invoices, tmp_path / "report.xlsx"))
    for name in ("Invoices", "Items"):
        sheet = workbook[name]
        assert sheet.page_setup.orientation == "landscape"
        assert sheet.page_setup.fitToWidth == 1
        assert sheet.sheet_properties.pageSetUpPr.fitToPage


def test_excel_dates_are_left_aligned(invoices, tmp_path):
    workbook = load_workbook(write_excel(invoices, tmp_path / "report.xlsx"))
    assert workbook["Invoices"]["C2"].alignment.horizontal == "left"


def test_excel_header_alignment_matches_columns(invoices, tmp_path):
    sheet = load_workbook(write_excel(invoices, tmp_path / "report.xlsx"))["Invoices"]
    assert sheet["H1"].alignment.horizontal == "right"  # Total
    assert sheet["I1"].alignment.horizontal == "center"  # Status
    assert sheet["I2"].alignment.horizontal == "center"
    assert sheet["A1"].alignment.horizontal == "left"
