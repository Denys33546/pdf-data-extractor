"""Write extracted invoices to Excel (.xlsx) and CSV files."""

import csv
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from pdf_extractor.extractor import ExtractedInvoice

INVOICE_COLUMNS = [
    "File",
    "Invoice No",
    "Date",
    "Vendor",
    "Customer",
    "Subtotal",
    "Tax",
    "Total",
    "Status",
    "Warnings",
]
ITEM_COLUMNS = [
    "File",
    "Invoice No",
    "Date",
    "Vendor",
    "Description",
    "Qty",
    "Unit Price",
    "Amount",
]
UNKNOWN_VENDOR = "(unknown)"

MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "yyyy-mm-dd"
HEADER_FILL = PatternFill("solid", fgColor="0F766E")
REVIEW_FILL = PatternFill("solid", fgColor="FDE2E2")
FORMULA_STARTS = ("=", "+", "-", "@")


def invoice_rows(invoices: list[ExtractedInvoice]) -> list[list]:
    """One row per invoice."""
    return [
        [
            inv.source,
            inv.number,
            inv.issue_date,
            inv.vendor,
            inv.customer,
            inv.subtotal,
            inv.tax,
            inv.total,
            "OK" if inv.is_valid else "CHECK",
            "; ".join(inv.warnings),
        ]
        for inv in invoices
    ]


def item_rows(invoices: list[ExtractedInvoice]) -> list[list]:
    """One row per line item, with the invoice it belongs to."""
    return [
        [
            inv.source,
            inv.number,
            inv.issue_date,
            inv.vendor,
            item.description,
            item.quantity,
            item.unit_price,
            item.amount,
        ]
        for inv in invoices
        for item in inv.items
    ]


def vendor_totals(invoices: list[ExtractedInvoice]) -> list[tuple[str, int, Decimal]]:
    """(vendor, number of invoices, total amount), biggest total first."""
    counts: dict[str, int] = defaultdict(int)
    totals: dict[str, Decimal] = defaultdict(lambda: Decimal("0.00"))
    for inv in invoices:
        vendor = inv.vendor or UNKNOWN_VENDOR
        counts[vendor] += 1
        totals[vendor] += inv.total or Decimal("0.00")
    return sorted(
        ((vendor, counts[vendor], totals[vendor]) for vendor in counts),
        key=lambda row: (-row[2], row[0]),
    )


def summary_metrics(invoices: list[ExtractedInvoice]) -> list[tuple[str, object]]:
    valid = sum(1 for inv in invoices if inv.is_valid)
    return [
        ("Invoices processed", len(invoices)),
        ("All checks passed", valid),
        ("Need review", len(invoices) - valid),
        ("Line items", sum(len(inv.items) for inv in invoices)),
        (
            "Total of all invoices",
            sum((inv.total or Decimal("0.00") for inv in invoices), Decimal("0.00")),
        ),
    ]


# ---------- Excel ----------


def _set_cell(ws: Worksheet, row: int, column: int, value: object) -> None:
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value.startswith(FORMULA_STARTS):
        cell.data_type = "s"  # keep text that looks like a formula as plain text


def _write_table(
    ws: Worksheet,
    columns: list[str],
    rows: list[list],
    money_columns: tuple[str, ...] = (),
) -> None:
    for column_number, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=column_number, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(
            horizontal=_alignment_for(name, money_columns), vertical="center"
        )
    for row_number, row in enumerate(rows, start=2):
        for column_number, value in enumerate(row, start=1):
            _set_cell(ws, row_number, column_number, value)
            cell = ws.cell(row=row_number, column=column_number)
            name = columns[column_number - 1]
            if name in money_columns:
                cell.number_format = MONEY_FORMAT
            elif name == "Date":
                cell.number_format = DATE_FORMAT
            if name == "Status":
                cell.alignment = Alignment(horizontal="center")
            elif name == "Date":
                cell.alignment = Alignment(horizontal="left")
    _fit_columns(ws, columns, rows)
    _print_setup(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(rows) + 1, 1)}"


def _alignment_for(column: str, money_columns: tuple[str, ...]) -> str:
    """Header alignment that matches the cells below it."""
    if column in money_columns:
        return "right"
    return "center" if column == "Status" else "left"


def _print_setup(ws: Worksheet) -> None:
    """Landscape, one page wide: the table does not get cut when printed or saved as PDF."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def _fit_columns(ws: Worksheet, columns: list[str], rows: list[list]) -> None:
    for column_number, name in enumerate(columns, start=1):
        longest = max([len(name)] + [len(str(row[column_number - 1] or "")) for row in rows])
        ws.column_dimensions[get_column_letter(column_number)].width = min(max(longest + 2, 10), 60)


def write_excel(invoices: list[ExtractedInvoice], path: Path) -> Path:
    """Create an .xlsx file with Summary, Invoices and Items sheets."""
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Summary"
    for row_number, (label, value) in enumerate(summary_metrics(invoices), start=1):
        summary.cell(row=row_number, column=1, value=label).font = Font(bold=True)
        cell = summary.cell(row=row_number, column=2, value=value)
        if isinstance(value, Decimal):
            cell.number_format = MONEY_FORMAT
    start = len(summary_metrics(invoices)) + 2
    for column_number, name in enumerate(["Vendor", "Invoices", "Total"], start=1):
        cell = summary.cell(row=start, column=column_number, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for offset, (vendor, count, total) in enumerate(vendor_totals(invoices), start=1):
        _set_cell(summary, start + offset, 1, vendor)
        summary.cell(row=start + offset, column=2, value=count)
        summary.cell(row=start + offset, column=3, value=total).number_format = MONEY_FORMAT
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 16
    summary.column_dimensions["C"].width = 16

    invoices_sheet = workbook.create_sheet("Invoices")
    rows = invoice_rows(invoices)
    _write_table(invoices_sheet, INVOICE_COLUMNS, rows, ("Subtotal", "Tax", "Total"))
    for row_number, invoice in enumerate(invoices, start=2):
        if not invoice.is_valid:
            for column_number in range(1, len(INVOICE_COLUMNS) + 1):
                invoices_sheet.cell(row=row_number, column=column_number).fill = REVIEW_FILL

    items_sheet = workbook.create_sheet("Items")
    _write_table(items_sheet, ITEM_COLUMNS, item_rows(invoices), ("Unit Price", "Amount"))

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


# ---------- CSV ----------


def _csv_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    if isinstance(value, str) and value.startswith(FORMULA_STARTS):
        return "'" + value  # spreadsheet apps would run it as a formula otherwise
    return str(value)  # dates become YYYY-MM-DD


def _write_csv_file(path: Path, columns: list[str], rows: list[list]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:  # BOM: Excel reads UTF-8
        writer = csv.writer(handle)
        writer.writerow(columns)
        writer.writerows([[_csv_value(value) for value in row] for row in rows])


def write_csv(invoices: list[ExtractedInvoice], path: Path) -> list[Path]:
    """Write `<name>_invoices.csv` and `<name>_items.csv` next to `path`."""
    path.parent.mkdir(parents=True, exist_ok=True)
    invoices_path = path.with_name(f"{path.stem}_invoices.csv")
    items_path = path.with_name(f"{path.stem}_items.csv")
    _write_csv_file(invoices_path, INVOICE_COLUMNS, invoice_rows(invoices))
    _write_csv_file(items_path, ITEM_COLUMNS, item_rows(invoices))
    return [invoices_path, items_path]
